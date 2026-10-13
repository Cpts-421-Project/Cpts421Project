"""
Bulk import: reads an Excel workbook (see bulk_upload_template.xlsx) and hands
each row to the existing ObjectService.create_object(), so it uses the same
code path as the Create Object form.

One row = one object. Columns (header matching ignores '*' and '(...)' hints):
  Image Source | Object Type* | Material | Display Date* | Start Date Range* |
  End Date Range* | Description | Inventory Number | Latitude | Longitude | Image URL

- Image URL: one URL, or several separated by semicolons / new lines.
- Image Source: Chicago-style citation, applied to every image in that row.
  It is matched to an existing `sources` row or created, then linked to the
  object's images after create_object runs (the Create form has no source field).

The whole file is validated first; nothing is created if there are errors.
Objects are then created one at a time (create_object commits each), so a
mid-run failure leaves earlier objects created and is reported per row.
"""
import io
import re
from typing import Any, Dict, List, Optional, Tuple

from openpyxl import load_workbook
from sqlalchemy.orm import Session

from ..models import Object, Image, Source
from ..logger import app_logger
from ..exceptions import BaseAppException
from .object_service import ObjectService

COLUMNS = {
    "image source": "citation", "object type": "object_type", "material": "material",
    "display date": "date_display", "start date range": "date_start",
    "end date range": "date_end", "description": "description",
    "inventory number": "inventory_number", "latitude": "latitude",
    "longitude": "longitude", "image url": "image_urls",
}
REQUIRED = {"object_type", "date_display", "date_start", "date_end"}
PLACEHOLDERS = {"", "yet to add", "n/a", "na", "?", "-", "none", "tbd"}
DEFAULT_IMAGE_TYPE, DEFAULT_VIEW_TYPE = "Photograph", "Front"  # same as the Create form


def _s(v: Any) -> Optional[str]:
    if v is None:
        return None
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    s = re.sub(r"[ \t]+", " ", str(v)).strip()
    return None if s.lower() in PLACEHOLDERS else s


def _num(v: Any) -> Optional[float]:
    s = _s(v)
    return float(s.replace(",", ".")) if s is not None else None


def _int(v: Any) -> Optional[int]:
    n = _num(v)
    if n is None:
        return None
    if n != int(n):
        raise ValueError
    return int(n)


def _norm_header(h: Any) -> str:
    return re.sub(r"\s+", " ", re.sub(r"\*|\(.*?\)", "", str(h or ""))).strip().lower()


def parse_workbook(file_bytes: bytes) -> Tuple[List[dict], List[dict], List[dict]]:
    errors: List[dict] = []
    warnings: List[dict] = []
    try:
        wb = load_workbook(io.BytesIO(file_bytes), read_only=True, data_only=True)
    except Exception:
        raise BaseAppException("Could not read file. Please upload a valid .xlsx workbook.", 400)

    ws = wb.worksheets[0]
    rows = list(ws.iter_rows(values_only=True))
    if len(rows) < 2:
        raise BaseAppException("The spreadsheet has no data rows.", 400)

    idx: Dict[int, str] = {}
    for i, h in enumerate(rows[0]):
        key = COLUMNS.get(_norm_header(h))
        if key and key not in idx.values():
            idx[i] = key
    missing = REQUIRED - set(idx.values())
    if missing:
        raise BaseAppException(
            f"Missing required column(s): {', '.join(sorted(missing))}. "
            "Headers must be in the first row of the first sheet.", 400)

    objects: List[dict] = []
    seen_inv: Dict[str, int] = {}
    for n, raw in enumerate(rows[1:], start=2):
        r = {k: (raw[i] if i < len(raw) else None) for i, k in idx.items()}
        if all(v is None or str(v).strip() == "" for v in r.values()):
            continue
        bad = lambda m, n=n: errors.append({"sheet": ws.title, "row": n, "message": m})

        rec: Dict[str, Any] = {k: _s(r.get(k)) for k in
                               ("object_type", "material", "date_display", "description", "inventory_number")}
        rec["findspot"] = None
        for k in ("date_start", "date_end"):
            try:
                rec[k] = _int(r.get(k))
            except ValueError:
                bad(f"{k.replace('_', ' ')} must be a whole number (negative for BC)."); rec[k] = None
        for k in ("latitude", "longitude"):
            try:
                rec[k] = _num(r.get(k))
            except ValueError:
                bad(f"{k} must be a number."); rec[k] = None

        for k, label in (("object_type", "Object Type"), ("date_display", "Display Date"),
                         ("date_start", "Start Date Range"), ("date_end", "End Date Range")):
            if rec[k] is None and not any(e["row"] == n and label.lower() in e["message"].lower()
                                          or e["row"] == n and k.replace("_", " ") in e["message"] for e in errors):
                bad(f"{label} is required.")
        if None not in (rec["date_start"], rec["date_end"]) and rec["date_start"] > rec["date_end"]:
            bad("Start Date Range is after End Date Range.")
        if rec["latitude"] is not None and not -90 <= rec["latitude"] <= 90:
            bad("Latitude must be between -90 and 90.")
        if rec["longitude"] is not None and not -180 <= rec["longitude"] <= 180:
            bad("Longitude must be between -180 and 180.")
        for k, lim in (("object_type", 255), ("material", 255), ("date_display", 255), ("inventory_number", 100)):
            if rec[k] and len(rec[k]) > lim:
                bad(f"{k.replace('_', ' ')} exceeds {lim} characters.")
        inv = rec["inventory_number"]
        if inv:
            if inv in seen_inv:
                bad(f"Inventory Number {inv} is also used on row {seen_inv[inv]}.")
            seen_inv[inv] = n

        urls = [u.strip() for u in re.split(r"[;\n]+", str(r.get("image_urls") or "")) if _s(u)]
        rec["images"] = [{"file_url": u, "image_type": DEFAULT_IMAGE_TYPE,
                          "view_type": DEFAULT_VIEW_TYPE} for u in urls]
        citation = _s(r.get("citation"))
        if citation and not urls:
            warnings.append({"sheet": ws.title, "row": n, "message": "Image Source given but no Image URL; source ignored."})
        if not urls:
            warnings.append({"sheet": ws.title, "row": n, "message": "No Image URL; object will have no images."})

        objects.append({"row": n, "citation": citation if urls else None, "data": rec})
    return objects, errors, warnings


def _link_source(db: Session, object_id: int, citation: str) -> None:
    src = db.query(Source).filter(Source.citation_text == citation).first()
    if not src:
        src = Source(citation_text=citation)
        db.add(src)
        db.flush()
    db.query(Image).filter(Image.object_id == object_id).update({"source_id": src.id})
    db.commit()


def import_objects_from_excel(db: Session, file_bytes: bytes, dry_run: bool = True) -> dict:
    objects, errors, warnings = parse_workbook(file_bytes)

    inv = [o["data"]["inventory_number"] for o in objects if o["data"]["inventory_number"]]
    if inv:
        existing = {r[0] for r in db.query(Object.inventory_number)
                    .filter(Object.inventory_number.in_(inv)).all()}
        for o in objects:
            if o["data"]["inventory_number"] in existing:
                errors.append({"sheet": "Sheet", "row": o["row"], "message":
                    f"Inventory Number {o['data']['inventory_number']} already exists in the database."})

    citations = {o["citation"] for o in objects if o["citation"]}
    known = {s.citation_text for s in db.query(Source)
             .filter(Source.citation_text.in_(citations)).all()} if citations else set()

    report = {
        "dry_run": dry_run,
        "objects_count": len(objects),
        "images_count": sum(len(o["data"]["images"]) for o in objects),
        "new_sources_count": len(citations - known),
        "errors": errors, "warnings": warnings,
        "preview": [{"key": o["row"], **{k: o["data"][k] for k in
                     ("object_type", "material", "date_display", "date_start", "date_end", "inventory_number")},
                     "image_count": len(o["data"]["images"])} for o in objects[:100]],
        "created_ids": [], "failed": [],
    }
    if dry_run or errors:
        return report

    service = ObjectService(db)
    for o in objects:
        try:
            created = service.create_object(dict(o["data"], images=list(o["data"]["images"])))
            if o["citation"]:
                _link_source(db, created["id"], o["citation"])
            report["created_ids"].append(created["id"])
        except Exception as exc:  # keep going; report per-row failures
            db.rollback()
            app_logger.error("Bulk import: row failed", exc_info=True, extra={"row": o["row"]})
            report["failed"].append({"row": o["row"],
                                     "message": str(getattr(exc, "message", None) or exc)})
    report["created_count"] = len(report["created_ids"])
    return report